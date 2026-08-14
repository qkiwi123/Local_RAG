"""DOCX adapter: python-docx -> ParsedDocument。

DOCX 没有稳定的物理页概念，因此每 30 个非空文本段落形成一个逻辑页；
标题、普通段落、列表和表格仍按文档 XML 阅读顺序保留。

C 层（OCR_EMBEDDED_IMAGES）：段落内嵌图片按阅读顺序原位插入识别文本块；
视觉产出带 ocr 属性标记，失败/超预算只产生 warning。
"""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from backend.config import settings
from backend.models.document_structure import (
    ParsedDocument,
    ParsedUnit,
    ParseWarning,
    SourceLocator,
)
from backend.services.parsers.assembler import BlockDraft, assemble_unit_from_drafts
from backend.services.parsers.ocr import (
    file_sha256,
    image_blob_sha256,
    ocr_attributes,
    run_vision,
)
from backend.services.parsers.office_zip import validate_office_package
from backend.services.parsers.tables import render_table_markdown
from backend.services.parsers.text import normalize_text
from backend.services.vision import normalize_image_bytes, recognize_text
from backend.services.vision_cache import VisionBudget


PARAGRAPHS_PER_VIRTUAL_PAGE = 30
_HEADING_STYLE_PATTERN = re.compile(r"(?:heading|标题)\s*([1-6])", re.IGNORECASE)
_LIST_STYLE_LEVEL_PATTERN = re.compile(r"(\d+)\s*$")


def _heading_level(paragraph: Paragraph) -> int | None:
    """优先读取 outline level，回退到 Heading/标题样式名。"""
    properties = paragraph._p.pPr
    outline_level = getattr(properties, "outlineLvl", None) if properties is not None else None
    outline_value = getattr(outline_level, "val", None)
    if outline_value is not None:
        try:
            return int(outline_value) + 1
        except (TypeError, ValueError):
            pass

    style_name = paragraph.style.name if paragraph.style else ""
    matched = _HEADING_STYLE_PATTERN.search(style_name)
    return int(matched.group(1)) if matched else None


def _list_level(paragraph: Paragraph) -> int | None:
    """读取直接编号层级；样式继承编号时回退到 List/列表样式名。"""
    properties = paragraph._p.pPr
    numbering = getattr(properties, "numPr", None) if properties is not None else None
    indentation = getattr(numbering, "ilvl", None) if numbering is not None else None
    indentation_value = getattr(indentation, "val", None)
    if indentation_value is not None:
        try:
            return int(indentation_value)
        except (TypeError, ValueError):
            return 0

    style_name = paragraph.style.name if paragraph.style else ""
    if "list" not in style_name.lower() and "列表" not in style_name:
        return None
    matched = _LIST_STYLE_LEVEL_PATTERN.search(style_name)
    return max(int(matched.group(1)) - 1, 0) if matched else 0


def _updated_heading_path(
    heading_path: list[str],
    heading_level: int,
    text: str,
) -> list[str]:
    """按标题级别更新路径，并允许文档直接从较深标题开始。"""
    return [*heading_path[: max(heading_level - 1, 0)], text]


def _table_rows(table: Table) -> list[list[str]]:
    """提取二维单元格，去除合并单元格重复值、空行和重复表头。"""
    rows: list[list[str]] = []
    for row in table.rows:
        seen_cells: set[int] = set()
        values: list[str] = []
        for cell in row.cells:
            cell_key = id(cell._tc)
            if cell_key in seen_cells:
                values.append("")
                continue
            seen_cells.add(cell_key)
            values.append(normalize_text(cell.text))

        if not any(values):
            continue
        if rows and values == rows[0]:
            continue
        rows.append(values)

    if not rows:
        return []
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def _paragraph_image_blobs(paragraph: Paragraph) -> list[tuple[str, bytes]]:
    """提取段落内嵌图片的关系 ID 与二进制内容（按出现顺序）。"""
    blips = paragraph._p.xpath(".//*[local-name()='blip']")
    result: list[tuple[str, bytes]] = []
    part = paragraph.part
    for blip in blips:
        relationship_id = blip.get(qn("r:embed"))
        if not relationship_id:
            continue
        try:
            blob = part.related_parts[relationship_id].blob
        except KeyError:
            continue
        if blob:
            result.append((relationship_id, blob))
    return result


def _embedded_ocr_enabled() -> bool:
    return settings.OCR_ENABLED and settings.OCR_EMBEDDED_IMAGES


def parse_docx_document(
    file_path: str,
    *,
    document_id: str,
    document_name: str | None = None,
) -> ParsedDocument:
    """按 XML 阅读顺序解析 DOCX，可选识别内嵌图片，输出 ParsedDocument。"""
    name = document_name or Path(file_path).name
    validate_office_package(file_path)
    document = Document(file_path)
    # 仅 C 层开启时才需要文件哈希（缓存键）。
    file_hash = file_sha256(file_path) if _embedded_ocr_enabled() else ""
    budget = VisionBudget()
    units: list[ParsedUnit] = []
    warnings: list[ParseWarning] = []

    virtual_page = 1
    paragraph_count = 0
    drafts: list[BlockDraft] = []
    heading_path: list[str] = []
    unit_heading_path: list[str] | None = None

    def warning_locator() -> SourceLocator:
        return SourceLocator(
            kind="virtual_page",
            unit_index=virtual_page - 1,
            location_label=f"逻辑页 {virtual_page}",
            page_number=virtual_page,
            heading_path=list(heading_path),
        )

    def append_draft(draft: BlockDraft, *, counts_as_paragraph: bool) -> None:
        nonlocal paragraph_count, unit_heading_path
        if unit_heading_path is None:
            unit_heading_path = list(heading_path)
        drafts.append(draft)
        if counts_as_paragraph:
            paragraph_count += 1

    def flush() -> None:
        nonlocal drafts, paragraph_count, virtual_page, unit_heading_path
        if not drafts:
            return
        locator = SourceLocator(
            kind="virtual_page",
            unit_index=virtual_page - 1,
            location_label=f"逻辑页 {virtual_page}",
            page_number=virtual_page,
            heading_path=list(unit_heading_path or heading_path),
        )
        units.append(
            assemble_unit_from_drafts(
                document_id=document_id,
                document_name=name,
                locator=locator,
                drafts=drafts,
            )
        )
        drafts = []
        paragraph_count = 0
        unit_heading_path = None
        virtual_page += 1

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = normalize_text(item.text)
            image_blobs = (
                _paragraph_image_blobs(item) if _embedded_ocr_enabled() else []
            )
            if not text and not image_blobs:
                continue

            # 只在下一段文本真正到来时分页。这样第 30 段后的表格仍属于
            # 当前逻辑页，而第 31 段从下一逻辑页开始。
            if paragraph_count >= PARAGRAPHS_PER_VIRTUAL_PAGE and text:
                flush()

            if text:
                style_name = item.style.name if item.style else ""
                heading_level = _heading_level(item)
                if heading_level is not None:
                    heading_path = _updated_heading_path(
                        heading_path,
                        heading_level,
                        text,
                    )
                    append_draft(
                        BlockDraft(
                            block_type="heading",
                            text=text,
                            heading_level=heading_level,
                            attributes={"style_name": style_name},
                        ),
                        counts_as_paragraph=True,
                    )
                else:
                    list_level = _list_level(item)
                    if list_level is not None:
                        append_draft(
                            BlockDraft(
                                block_type="list_item",
                                text=text,
                                list_level=list_level,
                                attributes={"style_name": style_name},
                            ),
                            counts_as_paragraph=True,
                        )
                    else:
                        append_draft(
                            BlockDraft(
                                block_type="paragraph",
                                text=text,
                                attributes={"style_name": style_name},
                            ),
                            counts_as_paragraph=True,
                        )

            for _relationship_id, blob in image_blobs:
                # EMF/WMF 等视觉模型不支持的格式先规范化；失败直接跳过。
                normalized = normalize_image_bytes(blob)
                if not normalized:
                    continue
                recognized = run_vision(
                    file_hash,
                    f"docx-img:{image_blob_sha256(blob)[:16]}",
                    normalized,
                    recognize_text,
                    budget=budget,
                    warnings=warnings,
                    locator=warning_locator(),
                    code="ocr_embedded_image_failed",
                )
                if recognized:
                    append_draft(
                        BlockDraft(
                            block_type="paragraph",
                            text=recognized,
                            attributes=ocr_attributes("embedded_image"),
                        ),
                        counts_as_paragraph=False,
                    )

        elif isinstance(item, Table):
            rows = _table_rows(item)
            if rows:
                append_draft(
                    BlockDraft(
                        block_type="table",
                        text=render_table_markdown(rows),
                        table_rows=rows,
                    ),
                    counts_as_paragraph=False,
                )

    flush()
    return ParsedDocument(units=units, warnings=warnings)


def parse_docx_units(
    file_path: str,
    *,
    document_id: str,
    document_name: str | None = None,
) -> list[ParsedUnit]:
    """统一 adapter 兼容入口；需要 warning 时使用 parse_docx_document。"""
    return parse_docx_document(
        file_path,
        document_id=document_id,
        document_name=document_name,
    ).units
