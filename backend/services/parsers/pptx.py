"""PPTX adapter: python-pptx -> ParsedDocument。

C 层（OCR_EMBEDDED_IMAGES）：图片 shape 按阅读顺序原位插入识别文本块；
视觉产出带 ocr 属性标记，失败/超预算只产生 warning。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from backend.config import settings
from backend.models.document_structure import (
    ParsedDocument,
    ParsedUnit,
    ParseWarning,
    SourceLocator,
)
from backend.services.parsers.assembler import BlockDraft, assemble_unit_from_drafts
from backend.services.parsers.ocr import (
    file_sha256,
    image_blob_sha256,
    ocr_attributes,
    run_vision,
)
from backend.services.parsers.office_zip import validate_office_package
from backend.services.parsers.tables import render_table_markdown
from backend.services.parsers.text import normalize_text
from backend.services.vision import recognize_text
from backend.services.vision_cache import VisionBudget


_LIST_PLACEHOLDER_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
}


def _shape_position(shape, original_index: int) -> tuple[int, int, int]:
    return (
        int(getattr(shape, "top", 0) or 0),
        int(getattr(shape, "left", 0) or 0),
        original_index,
    )


def _has_direct_bullet(paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is None:
        return False
    child_names = {
        child.tag.rsplit("}", 1)[-1]
        for child in properties
    }
    if "buNone" in child_names:
        return False
    return bool(child_names & {"buChar", "buAutoNum", "buBlip"})


def _is_list_paragraph(shape, paragraph) -> bool:
    if paragraph.level > 0 or _has_direct_bullet(paragraph):
        return True
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in _LIST_PLACEHOLDER_TYPES
    except (AttributeError, ValueError):
        return False


def _table_rows(table) -> list[list[str]]:
    rows = [
        [normalize_text(cell.text) for cell in row.cells]
        for row in table.rows
    ]
    return [row for row in rows if any(row)]


def _embedded_ocr_enabled() -> bool:
    return settings.OCR_ENABLED and settings.OCR_EMBEDDED_IMAGES


@dataclass
class _SlideVisionContext:
    """单张幻灯片内的视觉识别上下文（预算与 warning 通道）。"""

    file_hash: str
    budget: VisionBudget
    warnings: list[ParseWarning]
    locator: SourceLocator


def _append_shape_drafts(
    shape,
    drafts: list[BlockDraft],
    *,
    group_path: tuple[str, ...] = (),
    vision: _SlideVisionContext | None = None,
) -> None:
    """递归提取 shape；PPTX 的关键正文经常位于组合形状中。"""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        child_path = (*group_path, shape.name)
        positioned_children = sorted(
            enumerate(shape.shapes),
            key=lambda item: _shape_position(item[1], item[0]),
        )
        for _index, child in positioned_children:
            _append_shape_drafts(
                child,
                drafts,
                group_path=child_path,
                vision=vision,
            )
        return

    attributes: dict[str, str | int | bool] = {"shape_name": shape.name}
    if group_path:
        attributes["group_path"] = " / ".join(group_path)

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        if vision is None or not _embedded_ocr_enabled():
            return
        try:
            blob = shape.image.blob
        except (AttributeError, ValueError):
            return
        if not blob:
            return
        recognized = run_vision(
            vision.file_hash,
            f"pptx-img:{image_blob_sha256(blob)[:16]}",
            blob,
            recognize_text,
            budget=vision.budget,
            warnings=vision.warnings,
            locator=vision.locator,
            code="ocr_embedded_image_failed",
        )
        if recognized:
            drafts.append(
                BlockDraft(
                    block_type="paragraph",
                    text=recognized,
                    attributes={
                        **attributes,
                        **ocr_attributes("embedded_image"),
                    },
                )
            )
        return

    if getattr(shape, "has_table", False):
        rows = _table_rows(shape.table)
        if rows:
            drafts.append(
                BlockDraft(
                    block_type="table",
                    text=render_table_markdown(rows),
                    table_rows=rows,
                    attributes=attributes,
                )
            )
        return

    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        if _is_list_paragraph(shape, paragraph):
            drafts.append(
                BlockDraft(
                    block_type="list_item",
                    text=text,
                    list_level=paragraph.level,
                    attributes=attributes,
                )
            )
        else:
            drafts.append(
                BlockDraft(
                    block_type="paragraph",
                    text=text,
                    attributes=attributes,
                )
            )


def parse_pptx_document(
    file_path: str,
    *,
    document_id: str,
    document_name: str | None = None,
) -> ParsedDocument:
    """解析原生文本和表格；可选识别内嵌图片；无原生文字 slide 只产生 warning。"""
    name = document_name or Path(file_path).name
    validate_office_package(file_path)
    presentation = Presentation(file_path)
    # 仅 C 层开启时才需要文件哈希（缓存键）。
    file_hash = file_sha256(file_path) if _embedded_ocr_enabled() else ""
    budget = VisionBudget()
    units: list[ParsedUnit] = []
    warnings: list[ParseWarning] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        locator = SourceLocator(
            kind="slide",
            unit_index=slide_number - 1,
            location_label=f"第 {slide_number} 张幻灯片",
            slide_number=slide_number,
        )
        drafts: list[BlockDraft] = []
        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None
        title_text = normalize_text(title_shape.text) if title_shape is not None else ""

        if title_text:
            drafts.append(
                BlockDraft(
                    block_type="heading",
                    text=title_text,
                    heading_level=1,
                    attributes={"shape_name": title_shape.name},
                )
            )

        vision = _SlideVisionContext(
            file_hash=file_hash,
            budget=budget,
            warnings=warnings,
            locator=locator,
        )
        positioned_shapes = sorted(
            [
                (shape, index)
                for index, shape in enumerate(slide.shapes)
                if shape.shape_id != title_id
            ],
            key=lambda item: _shape_position(item[0], item[1]),
        )
        for shape, _original_index in positioned_shapes:
            _append_shape_drafts(shape, drafts, vision=vision)

        if not drafts:
            warnings.append(
                ParseWarning(
                    code="pptx_no_native_text",
                    message=(
                        f"第 {slide_number} 张幻灯片未提取到原生文本或表格"
                    ),
                    locator=locator,
                )
            )
            continue

        heading_path = [title_text] if title_text else []
        unit_locator = locator.model_copy(update={"heading_path": heading_path})
        units.append(
            assemble_unit_from_drafts(
                document_id=document_id,
                document_name=name,
                locator=unit_locator,
                drafts=drafts,
            )
        )

    return ParsedDocument(units=units, warnings=warnings)


def parse_pptx_units(
    file_path: str,
    *,
    document_id: str,
    document_name: str | None = None,
) -> list[ParsedUnit]:
    """统一 adapter 兼容入口；需要 warning 时使用 parse_pptx_document。"""
    return parse_pptx_document(
        file_path,
        document_id=document_id,
        document_name=document_name,
    ).units
