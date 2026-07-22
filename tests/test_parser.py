"""
测试 backend/services/parser.py — 文档解析器。

覆盖：
- PDF 解析：正确提取页数与页码
- TXT 解析：正确返回文本
- DOCX 解析：正确返回文本
- 不支持的文件类型抛出 ValueError
- 空 PDF 页面标记处理
"""

import pytest
from docx import Document

from backend.services.parser import parse_document


class TestParseDocument:
    """解析器核心功能测试。"""

    def test_parse_pdf_pages(self):
        """PDF 解析后得到的页数正确，页码连续递增。"""
        pages = parse_document("tests/fixtures/sample.pdf")
        assert len(pages) == 3, f"Expected 3 pages, got {len(pages)}"
        assert [p["page_number"] for p in pages] == [1, 2, 3]
        assert all(p["document_name"] == "sample.pdf" for p in pages)

    def test_parse_txt_returns_content(self):
        """TXT 解析能正常返回至少 1 页内容。"""
        pages = parse_document("tests/fixtures/sample.txt")
        assert len(pages) >= 1
        assert len(pages[0]["text"]) > 0
        assert pages[0]["page_number"] == 1

    def test_parse_docx_returns_content(self):
        """DOCX 解析能正常返回至少 1 页内容。"""
        pages = parse_document("tests/fixtures/sample.docx")
        assert len(pages) >= 1
        assert len(pages[0]["text"]) > 0
        assert pages[0]["page_number"] == 1

    def test_unsupported_format_raises_error(self):
        """不支持的文件类型抛出 ValueError。"""
        with pytest.raises(ValueError, match="不支持的文件格式"):
            parse_document("tests/fixtures/sample.png")

    def test_pdf_page_content_not_empty(self):
        """PDF 每页内容不应为空。"""
        pages = parse_document("tests/fixtures/sample.pdf")
        for page in pages:
            assert len(page["text"]) > 0, f"Page {page['page_number']} is empty"
            # 确保不是空页面占位标记
            assert "无可提取文本" not in page["text"]

    def test_custom_document_name(self):
        """传入自定义文档名时应覆盖从路径提取的文件名。"""
        pages = parse_document(
            "tests/fixtures/sample.txt",
            document_name="custom_name.txt"
        )
        assert all(p["document_name"] == "custom_name.txt" for p in pages)

    def test_document_id_is_preserved(self):
        pages = parse_document(
            "tests/fixtures/sample.txt",
            document_name="custom_name.txt",
            document_id="stable-doc-id",
        )
        assert all(p["document_id"] == "stable-doc-id" for p in pages)

    def test_text_normalization_is_conservative(self, tmp_path):
        source = tmp_path / "whitespace.txt"
        source.write_text(
            "ＡＩ\u3000技术\t应用，保留中文括号（ＳＯＰ）"
            "\r\n\r\n\r\n第二段  保留数字 ２０２４-２０２５",
            encoding="utf-8",
        )

        pages = parse_document(str(source))

        assert pages[0]["text"] == (
            "AI 技术 应用，保留中文括号（SOP）\n\n第二段 保留数字 2024-2025"
        )
        assert pages[0]["content_type"] == "paragraph"

    def test_docx_preserves_paragraph_table_order(self, tmp_path):
        source = tmp_path / "ordered.docx"
        document = Document()
        document.add_paragraph("表格前的正文")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "技术"
        table.cell(0, 1).text = "成熟度"
        table.cell(1, 0).text = "RAG"
        table.cell(1, 1).text = "成熟"
        document.add_paragraph("表格后的正文")
        document.save(source)

        pages = parse_document(str(source))

        assert [page["content_type"] for page in pages] == [
            "paragraph",
            "table",
            "paragraph",
        ]
        assert pages[0]["text"] == "表格前的正文"
        assert pages[1]["text"] == (
            "| 技术 | 成熟度 |\n"
            "| --- | --- |\n"
            "| RAG | 成熟 |"
        )
        assert pages[2]["text"] == "表格后的正文"

    def test_docx_table_handles_merged_empty_and_repeated_header_rows(self, tmp_path):
        source = tmp_path / "table-cases.docx"
        document = Document()
        table = document.add_table(rows=5, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "模型"
        table.cell(1, 1).text = "BGE"
        table.cell(2, 0).text = "字段"
        table.cell(2, 1).text = "值"
        table.cell(4, 0).merge(table.cell(4, 1)).text = "合并说明"
        document.save(source)

        pages = parse_document(str(source))
        table_text = next(
            page["text"] for page in pages if page["content_type"] == "table"
        )

        assert table_text.count("| 字段 | 值 |") == 1
        assert "| 模型 | BGE |" in table_text
        assert "| 合并说明 |  |" in table_text
        assert "|  |  |" not in table_text

    def test_parser_version_is_injected_into_every_block(self):
        """每种格式的每个 parser block 都必须携带固定的 PARSER_VERSION。"""
        from backend.services.parser import PARSER_VERSION

        fixtures = [
            "tests/fixtures/sample.pdf",
            "tests/fixtures/sample.txt",
            "tests/fixtures/sample.docx",
        ]
        for fixture_path in fixtures:
            pages = parse_document(fixture_path)
            assert len(pages) > 0, f"{fixture_path} returned no pages"
            for page in pages:
                assert page.get("parser_version") == PARSER_VERSION, (
                    f"{fixture_path} page {page['page_number']} "
                    f"expected parser_version={PARSER_VERSION}, "
                    f"got {page.get('parser_version')}"
                )

    def test_docx_table_keeps_virtual_page_number(self, tmp_path):
        source = tmp_path / "page-number.docx"
        document = Document()
        for index in range(31):
            document.add_paragraph(f"第 {index + 1} 个段落")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "技术"
        table.cell(0, 1).text = "时间"
        table.cell(1, 0).text = "RAG"
        table.cell(1, 1).text = "2024-2025"
        document.save(source)

        pages = parse_document(str(source))
        table_page = next(page for page in pages if page["content_type"] == "table")

        assert table_page["page_number"] == 2
