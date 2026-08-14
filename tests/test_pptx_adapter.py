import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from backend.config import settings
from backend.services.parsers import PARSER_VERSION
from backend.services.parsers.pptx import parse_pptx_document


def test_pptx_preserves_title_positioned_text_lists_and_table(tmp_path):
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "检索评测"

    lower = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    lower.text_frame.paragraphs[0].text = "后出现但位置较下"
    upper = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    upper.text_frame.paragraphs[0].text = "后创建但位置较上"
    nested = upper.text_frame.add_paragraph()
    nested.text = "嵌套项"
    nested.level = 1

    table_shape = slide.shapes.add_table(
        2,
        2,
        Inches(1),
        Inches(4),
        Inches(5),
        Inches(1),
    )
    table_shape.table.cell(0, 0).text = "指标"
    table_shape.table.cell(0, 1).text = "结果"
    table_shape.table.cell(1, 0).text = "MRR"
    table_shape.table.cell(1, 1).text = "0.8"
    presentation.save(source)

    result = parse_pptx_document(
        str(source),
        document_id="doc-slides",
        document_name="课件.pptx",
    )

    assert result.warnings == []
    assert len(result.units) == 1
    unit = result.units[0]
    assert unit.parser_version == PARSER_VERSION
    assert unit.locator.kind == "slide"
    assert unit.locator.slide_number == 1
    assert unit.locator.location_label == "第 1 张幻灯片"
    assert unit.locator.heading_path == ["检索评测"]
    assert [block.block_type for block in unit.blocks] == [
        "heading",
        "paragraph",
        "list_item",
        "paragraph",
        "table",
    ]
    assert [block.text for block in unit.blocks[1:4]] == [
        "后创建但位置较上",
        "嵌套项",
        "后出现但位置较下",
    ]
    assert unit.blocks[2].list_level == 1
    assert unit.blocks[-1].table_rows == [
        ["指标", "结果"],
        ["MRR", "0.8"],
    ]


def test_pptx_extracts_text_from_nested_group_shapes(tmp_path):
    source = tmp_path / "grouped.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "组合形状"

    outer = slide.shapes.add_group_shape()
    lower = outer.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(1),
    )
    lower.text = "组合内下方正文"
    nested = outer.shapes.add_group_shape()
    upper = nested.shapes.add_textbox(
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(1),
    )
    upper.text = "嵌套组合正文"
    presentation.save(source)

    result = parse_pptx_document(str(source), document_id="doc-grouped")

    assert result.warnings == []
    unit = result.units[0]
    assert [block.text for block in unit.blocks] == [
        "组合形状",
        "嵌套组合正文",
        "组合内下方正文",
    ]
    assert unit.blocks[1].attributes["group_path"] == (
        f"{outer.name} / {nested.name}"
    )
    assert unit.blocks[2].attributes["group_path"] == outer.name


def test_pptx_image_only_slide_emits_warning_without_fabricated_unit(tmp_path):
    source = tmp_path / "image-only.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    # 生成一个最小图片后插入；解析器只应看到 picture shape 并跳过。
    from PIL import Image

    image_path = tmp_path / "pixel.png"
    Image.new("RGB", (1, 1), "white").save(image_path)
    picture = slide.shapes.add_picture(str(image_path), 0, 0)
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    presentation.save(source)

    result = parse_pptx_document(str(source), document_id="doc-image")

    assert result.units == []
    assert len(result.warnings) == 1
    warning = result.warnings[0]
    assert warning.code == "pptx_no_native_text"
    assert warning.locator.slide_number == 1
    assert "未提取到原生文本" in warning.message


def test_pptx_slide_unit_indexes_keep_original_gaps(tmp_path):
    source = tmp_path / "gaps.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "第二张有文字"
    presentation.save(source)

    result = parse_pptx_document(str(source), document_id="doc-gaps")

    assert len(result.warnings) == 1
    assert len(result.units) == 1
    assert result.units[0].locator.unit_index == 1
    assert result.units[0].locator.slide_number == 2


# ─── C 层：内嵌图片识别（mock 视觉调用） ───────────────────────────

def _image_pptx(path, image_path):
    """生成「标题 + 图片 + 正文」的 PPTX fixture。"""
    from PIL import Image

    Image.new("RGB", (60, 60), (0, 0, 255)).save(image_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "图片页"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1.5))
    textbox = slide.shapes.add_textbox(
        Inches(1), Inches(3), Inches(4), Inches(1)
    )
    textbox.text = "图片下方的正文"
    presentation.save(path)
    return path


@pytest.fixture
def embedded_ocr_enabled(monkeypatch, temp_dir):
    from backend.services.parsers.ocr import clear_failure_memo

    clear_failure_memo()  # 用例间隔离进程级失败记忆
    monkeypatch.setattr(settings, "OCR_ENABLED", True)
    monkeypatch.setattr(settings, "OCR_EMBEDDED_IMAGES", True)
    monkeypatch.setattr(settings, "VISION_CACHE_DIR", f"{temp_dir}/vision_cache")


def test_pptx_embedded_image_inserts_ocr_block_in_place(
    tmp_path, monkeypatch, embedded_ocr_enabled
):
    from backend.services.parsers import pptx as pptx_adapter

    source = _image_pptx(tmp_path / "with-image.pptx", tmp_path / "pixel.png")
    monkeypatch.setattr(
        pptx_adapter, "recognize_text", lambda blob: "图片中的文字"
    )

    result = parse_pptx_document(str(source), document_id="doc-img")

    assert result.warnings == []
    assert [block.text for block in result.units[0].blocks] == [
        "图片页",
        "图片中的文字",
        "图片下方的正文",
    ]
    ocr_block = result.units[0].blocks[1]
    assert ocr_block.attributes["ocr"] is True
    assert ocr_block.attributes["vision_kind"] == "embedded_image"
    assert ocr_block.attributes["vision_model"] == settings.VISION_MODEL
    assert ocr_block.attributes["shape_name"]


def test_pptx_embedded_ocr_disabled_skips_images(tmp_path):
    source = _image_pptx(tmp_path / "with-image.pptx", tmp_path / "pixel.png")

    result = parse_pptx_document(str(source), document_id="doc-off")

    assert result.warnings == []
    assert [block.text for block in result.units[0].blocks] == [
        "图片页",
        "图片下方的正文",
    ]


def test_pptx_embedded_ocr_failure_adds_warning(
    tmp_path, monkeypatch, embedded_ocr_enabled
):
    from backend.services.parsers import pptx as pptx_adapter
    from backend.services.vision import VisionError

    source = _image_pptx(tmp_path / "with-image.pptx", tmp_path / "pixel.png")
    monkeypatch.setattr(
        pptx_adapter,
        "recognize_text",
        lambda blob: (_ for _ in ()).throw(VisionError("识别失败")),
    )

    result = parse_pptx_document(str(source), document_id="doc-fail")

    assert len(result.warnings) == 1
    assert result.warnings[0].code == "ocr_embedded_image_failed"
    assert "识别失败" in result.warnings[0].message
    assert [block.text for block in result.units[0].blocks] == [
        "图片页",
        "图片下方的正文",
    ]
