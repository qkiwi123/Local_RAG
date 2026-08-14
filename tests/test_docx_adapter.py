import pytest
from docx import Document

from backend.config import settings
from backend.services.parsers import PARSER_VERSION
from backend.services.parsers.docx import (
    PARAGRAPHS_PER_VIRTUAL_PAGE,
    parse_docx_document,
    parse_docx_units,
)


def test_docx_preserves_heading_list_table_and_reading_order(tmp_path):
    source = tmp_path / "structured.docx"
    document = Document()
    document.add_heading("部署指南", level=1)
    document.add_paragraph("准备运行环境。")
    document.add_paragraph("安装依赖", style="List Bullet 2")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "组件"
    table.cell(0, 1).text = "版本"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "3.14"
    document.save(source)

    units = parse_docx_units(
        str(source),
        document_id="doc-1",
        document_name="指南.docx",
    )

    assert len(units) == 1
    unit = units[0]
    assert unit.parser_version == PARSER_VERSION
    assert unit.locator.kind == "virtual_page"
    assert unit.locator.location_label == "逻辑页 1"
    assert unit.locator.heading_path == ["部署指南"]
    assert [block.block_type for block in unit.blocks] == [
        "heading",
        "paragraph",
        "list_item",
        "table",
    ]
    assert unit.blocks[0].heading_level == 1
    assert unit.blocks[2].list_level == 1
    assert unit.blocks[3].table_rows == [
        ["组件", "版本"],
        ["Python", "3.14"],
    ]
    assert unit.canonical_text.endswith(
        "| 组件 | 版本 |\n| --- | --- |\n| Python | 3.14 |"
    )


def test_docx_virtual_page_boundary_keeps_table_and_heading_context(tmp_path):
    source = tmp_path / "pages.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    for index in range(PARAGRAPHS_PER_VIRTUAL_PAGE - 1):
        document.add_paragraph(f"第 {index + 1} 段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "页内"
    table.cell(0, 1).text = "表格"
    document.add_paragraph("下一逻辑页正文")
    document.save(source)

    units = parse_docx_units(str(source), document_id="doc-pages")

    assert len(units) == 2
    assert [unit.locator.page_number for unit in units] == [1, 2]
    assert units[0].blocks[-1].block_type == "table"
    assert units[1].canonical_text == "下一逻辑页正文"
    assert units[0].locator.heading_path == ["第一章"]
    assert units[1].locator.heading_path == ["第一章"]


def test_docx_table_removes_merged_duplicates_empty_rows_and_repeated_header(tmp_path):
    source = tmp_path / "table-cases.docx"
    document = Document()
    table = document.add_table(rows=5, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "模型"
    table.cell(1, 1).text = "BGE"
    table.cell(2, 0).text = "字段"
    table.cell(2, 1).text = "值"
    table.cell(4, 0).merge(table.cell(4, 1)).text = "合并说明"
    document.save(source)

    unit = parse_docx_units(str(source), document_id="doc-table")[0]
    block = unit.blocks[0]

    assert block.block_type == "table"
    assert block.table_rows == [
        ["字段", "值"],
        ["模型", "BGE"],
        ["合并说明", ""],
    ]
    assert block.text.count("| 字段 | 值 |") == 1


def test_docx_ids_are_deterministic(tmp_path):
    source = tmp_path / "stable.docx"
    document = Document()
    document.add_heading("标题", level=1)
    document.add_paragraph("正文")
    document.save(source)

    first = parse_docx_units(str(source), document_id="stable-doc")
    second = parse_docx_units(str(source), document_id="stable-doc")

    assert [unit.unit_id for unit in first] == [unit.unit_id for unit in second]
    assert [
        [block.block_id for block in unit.blocks]
        for unit in first
    ] == [
        [block.block_id for block in unit.blocks]
        for unit in second
    ]


# ─── C 层：内嵌图片识别（mock 视觉调用） ───────────────────────────

def _image_docx(path):
    """生成含「文字段落 + 内嵌图片」的 DOCX fixture。"""
    from io import BytesIO

    from PIL import Image

    image = Image.new("RGB", (60, 60), (0, 0, 255))
    buffer = BytesIO()
    image.save(buffer, format="JPEG")
    buffer.seek(0)

    document = Document()
    document.add_paragraph("图片前面的段落。")
    document.add_picture(buffer, width=2 * 914400 // 1440)
    document.add_paragraph("图片后面的段落。")
    document.save(path)
    return path


@pytest.fixture
def embedded_ocr_enabled(monkeypatch, temp_dir):
    monkeypatch.setattr(settings, "OCR_ENABLED", True)
    monkeypatch.setattr(settings, "OCR_EMBEDDED_IMAGES", True)
    monkeypatch.setattr(settings, "VISION_CACHE_DIR", f"{temp_dir}/vision_cache")


def test_docx_embedded_image_inserts_ocr_block_in_place(
    tmp_path, monkeypatch, embedded_ocr_enabled
):
    from backend.services.parsers import docx as docx_adapter

    source = _image_docx(tmp_path / "with-image.docx")
    monkeypatch.setattr(
        docx_adapter, "recognize_text", lambda blob: "图片中的文字"
    )

    result = parse_docx_document(str(source), document_id="doc-img")

    assert result.warnings == []
    assert [block.text for block in result.units[0].blocks] == [
        "图片前面的段落。",
        "图片中的文字",
        "图片后面的段落。",
    ]
    ocr_block = result.units[0].blocks[1]
    assert ocr_block.attributes == {
        "ocr": True,
        "vision_model": settings.VISION_MODEL,
        "vision_kind": "embedded_image",
    }


def test_docx_embedded_ocr_disabled_skips_images(tmp_path):
    source = _image_docx(tmp_path / "with-image.docx")

    result = parse_docx_document(str(source), document_id="doc-off")

    assert result.warnings == []
    assert [block.text for block in result.units[0].blocks] == [
        "图片前面的段落。",
        "图片后面的段落。",
    ]
    assert all(not block.attributes.get("ocr") for block in result.units[0].blocks)


def test_docx_embedded_ocr_failure_adds_warning(
    tmp_path, monkeypatch, embedded_ocr_enabled
):
    from backend.services.parsers import docx as docx_adapter
    from backend.services.vision import VisionError

    source = _image_docx(tmp_path / "with-image.docx")
    monkeypatch.setattr(
        docx_adapter,
        "recognize_text",
        lambda blob: (_ for _ in ()).throw(VisionError("识别失败")),
    )

    result = parse_docx_document(str(source), document_id="doc-fail")

    assert len(result.warnings) == 1
    assert result.warnings[0].code == "ocr_embedded_image_failed"
    assert "识别失败" in result.warnings[0].message
    assert [block.text for block in result.units[0].blocks] == [
        "图片前面的段落。",
        "图片后面的段落。",
    ]
