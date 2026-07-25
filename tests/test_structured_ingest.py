from pathlib import Path

from docx import Document
from pptx import Presentation

from backend.services.chunker import chunk_units
from backend.services.document_store import create_document
from backend.services.metadata_schema import bootstrap_structured_database
from backend.services.parsers import PARSER_VERSION, parse_document_structure
from backend.services.unit_store import list_document_units, replace_document_units


def _build_sources(tmp_path):
    markdown = tmp_path / "sample.md"
    markdown.write_text("# Markdown 标题\n\n正文内容。", encoding="utf-8")

    pptx = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX 标题"
    presentation.save(pptx)

    # 使用生成文件覆盖 DOCX 的结构化解析真实路径；PDF/TXT 使用已有金丝雀。
    docx = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("DOCX 标题", level=1)
    document.add_paragraph("正文内容。")
    document.save(docx)

    return [
        Path("tests/fixtures/sample.pdf"),
        docx,
        Path("tests/fixtures/sample.txt"),
        markdown,
        pptx,
    ]


def test_five_formats_flow_through_units_chunks_sqlite_and_chroma_metadata(
    tmp_path,
    monkeypatch,
):
    from backend.services import vector_store

    db_path = str(tmp_path / "metadata.db")
    bootstrap_structured_database(
        db_path,
        created_at="2026-07-26T00:00:00+00:00",
    )
    captured_batches = []

    class FakeCollection:
        def add(self, **kwargs):
            captured_batches.append(kwargs)

    monkeypatch.setattr(vector_store, "_collection", FakeCollection())
    monkeypatch.setattr(
        vector_store,
        "embed_texts",
        lambda texts: [[0.0] for _ in texts],
    )

    for index, source in enumerate(_build_sources(tmp_path), start=1):
        document_id = f"doc-{index}"
        create_document(
            document_id,
            source.name,
            f"{document_id}{source.suffix}",
            source.stat().st_size,
            db_path,
        )
        result = parse_document_structure(
            str(source),
            document_id=document_id,
        )
        assert result.units, source.suffix
        assert result.warnings == [], source.suffix

        chunks = chunk_units(
            result.units,
            chunk_size=400,
            overlap=50,
        )
        assert chunks, source.suffix
        for chunk in chunks:
            unit = next(
                unit
                for unit in result.units
                if unit.unit_id == chunk.metadata.unit_id
            )
            assert unit.canonical_text[
                chunk.metadata.start_char:chunk.metadata.end_char
            ] == chunk.content

        replace_document_units(
            document_id,
            PARSER_VERSION,
            result.units,
            db_path=db_path,
        )
        assert list_document_units(
            document_id,
            PARSER_VERSION,
            db_path=db_path,
        ) == result.units
        assert vector_store.add_chunks(chunks) == len(chunks)

    assert len(captured_batches) == 5
    for batch in captured_batches:
        assert batch["ids"]
        for metadata in batch["metadatas"]:
            assert metadata["parser_version"] == PARSER_VERSION
            assert metadata["unit_id"]
            assert metadata["location_label"]
            assert isinstance(metadata["heading_path"], str)
