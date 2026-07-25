import pytest

from pptx import Presentation

from backend.services.parsers import parse_document_structure, parse_document_units


def test_registry_dispatches_txt_and_preserves_identity(tmp_path):
    source = tmp_path / "sample.TXT"
    source.write_text("正文", encoding="utf-8")

    units = parse_document_units(
        str(source),
        document_id="doc-1",
        document_name="资料.txt",
    )

    assert len(units) == 1
    assert units[0].document_id == "doc-1"
    assert units[0].document_name == "资料.txt"


def test_registry_rejects_unsupported_suffix(tmp_path):
    source = tmp_path / "sample.xlsx"
    source.write_bytes(b"not-an-excel-file")

    with pytest.raises(ValueError, match=r"\.xlsx.*当前支持"):
        parse_document_units(str(source), document_id="doc-unsupported")


def test_registry_dispatches_markdown(tmp_path):
    source = tmp_path / "sample.md"
    source.write_text("# 标题\n\n正文。", encoding="utf-8")

    units = parse_document_units(str(source), document_id="doc-md")

    assert len(units) == 1
    assert units[0].locator.kind == "section"
    assert units[0].canonical_text == "标题\n\n正文。"


def test_registry_returns_pptx_warnings(tmp_path):
    source = tmp_path / "empty-slide.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)

    result = parse_document_structure(str(source), document_id="doc-pptx")

    assert result.units == []
    assert len(result.warnings) == 1
    assert result.warnings[0].locator.kind == "slide"
